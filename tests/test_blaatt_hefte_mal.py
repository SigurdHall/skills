"""Fase 1: kontrakten, eksempeldokumentet, malspesifikasjonen, malen og 2024-fasiten.

Offline. Skriptene importeres direkte fra skill-mappen med importlib.
"""
from __future__ import annotations

import copy
import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

SKILL = Path(__file__).resolve().parents[1] / "skills/knowledge-management/blatt-hefte"
SCRIPTS = SKILL / "scripts"
ASSETS = SKILL / "assets"


def _last(navn: str):
    spec = importlib.util.spec_from_file_location(navn, SCRIPTS / f"{navn}.py")
    modul = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modul)
    return modul


kontrakt = _last("kontrakt")
mal = _last("mal")


@pytest.fixture(scope="module")
def eksempel() -> dict:
    return json.loads((ASSETS / "eksempel-2025-forslag.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def spec() -> dict:
    return mal.last_spec(ASSETS / "mal-rammeark.json")


@pytest.fixture(scope="module")
def malbok():
    return openpyxl.load_workbook(ASSETS / "mal-rammeark.xlsx")


@pytest.fixture(scope="module")
def fasit() -> dict:
    return json.loads((ASSETS / "fasit-2024.json").read_text(encoding="utf-8"))


# --- kontrakten og eksempeldokumentet -------------------------------------------------

def test_eksempeldokumentet_er_gyldig(eksempel):
    assert kontrakt.valider(eksempel) == []


def test_eksempeldokumentet_har_uit_raden(eksempel):
    uit = next(r for r in eksempel["hovedtabell"]["rader"] if r["kortkode_norm"] == "UIT")
    assert uit["verdier"] == [4061059, 157470, -17034, -24255, -25307, 20487,
                              -17329, -5653, -1091, 20800, 520, -14214, 4155453]
    assert eksempel["prisjustering"]["sats_prosent"] == 3.8
    assert eksempel["prisjustering"]["ren_sats"] is False
    assert eksempel["satser"]["lukket_ramme"] is None
    assert sum(1 for i in eksempel["institusjoner"]["liste"] if i["universitet"]) == 10


def test_alle_rader_summerer(eksempel):
    for rad in eksempel["hovedtabell"]["rader"]:
        assert sum(v or 0 for v in rad["verdier"][:-1]) == rad["verdier"][-1], rad["kortkode"]


def test_manglende_felt_gir_feil(eksempel):
    d = copy.deepcopy(eksempel)
    del d["prisjustering"]
    feil = kontrakt.valider(d)
    assert any("prisjustering" in f for f in feil)


def test_feil_lengde_gir_feil(eksempel):
    d = copy.deepcopy(eksempel)
    d["hovedtabell"]["rader"][0]["verdier"] = d["hovedtabell"]["rader"][0]["verdier"][:-1]
    feil = kontrakt.valider(d)
    assert any("verdier mot" in f for f in feil)


def test_feil_sumindeks_gir_feil(eksempel):
    d = copy.deepcopy(eksempel)
    d["hovedtabell"]["sum_indeks"] = 3
    assert any("sum_indeks" in f for f in kontrakt.valider(d))


# --- malspesifikasjonen mot malen -----------------------------------------------------

def test_arkrekkefolge(malbok):
    assert malbok.sheetnames == ["Hovedtall", "Hovedposter", "Resultat",
                                 "Sektor", "Satser", "Kilder", "_meta"]


def test_hver_navngitt_celle_finnes(spec, malbok):
    layout = mal.legg_ut(spec)
    meta = {r[1]: (r[2], r[3]) for r in malbok["_meta"].iter_rows(min_row=2, values_only=True)
            if r[0] == "celle"}
    assert layout.adresser, "spesifikasjonen navngir ingen celler"
    for navn, (ark, celle) in layout.adresser.items():
        assert meta.get(navn) == (ark, celle), navn
        assert navn in malbok.defined_names, navn
        assert malbok[ark][celle] is not None


def test_hver_formeltekst_og_tallformat_finnes(spec, malbok):
    layout = mal.legg_ut(spec)
    formler = 0
    for o in layout.oppgaver:
        celle = malbok[o["ark"]][o["celle"]]
        if o.get("formel"):
            assert celle.value == o["formel"], o["celle"]
            formler += 1
        elif o.get("verdi") is not None:
            assert celle.value == o["verdi"], o["celle"]
        if o.get("format"):
            assert celle.number_format == spec["tallformat"][o["format"]], o["celle"]
    assert formler >= 20


def test_startceller_for_dynamiske_blokker(spec, malbok):
    layout = mal.legg_ut(spec)
    for navn in ["hovedposter_start", "resultat_start", "sektor_universiteter_start",
                 "sektor_alle_start", "satser_open_ramme_start", "satser_lukket_ramme_start"]:
        assert navn in layout.radmerker, navn
        assert navn in layout.adresser, navn


def test_bygg_mal_gir_identiske_celler(spec, malbok, tmp_path):
    ny = tmp_path / "mal.xlsx"
    mal.bygg_mal(spec, ny)
    bygget = openpyxl.load_workbook(ny)
    assert bygget.sheetnames == malbok.sheetnames
    for navn in malbok.sheetnames:
        a, b = malbok[navn], bygget[navn]
        assert (a.max_row, a.max_column) == (b.max_row, b.max_column), navn
        for rad_a, rad_b in zip(a.iter_rows(), b.iter_rows()):
            for ca, cb in zip(rad_a, rad_b):
                assert (ca.coordinate, ca.value, ca.number_format) == \
                       (cb.coordinate, cb.value, cb.number_format), f"{navn}!{ca.coordinate}"


def test_radantall_flytter_blokkene(spec):
    standard = mal.legg_ut(spec)
    stor = mal.legg_ut(spec, radantall={"hovedposter": 13, "sektor_alle": 25})
    assert stor.radmerker["sektor_alle_slutt"] - stor.radmerker["sektor_alle_start"] == 24
    assert standard.radmerker["hovedposter_start"] == stor.radmerker["hovedposter_start"]


def test_uit_stil(spec):
    assert spec["stil"]["overskrift_bakgrunn"].endswith("003349")
    assert spec["stil"]["overskrift_tekst"].endswith("FFFFFF")
    assert spec["stil"]["vis_rutenett"] is False


def test_rutenett_av_i_malen(malbok):
    for navn in malbok.sheetnames:
        if navn != "_meta":
            assert malbok[navn].sheet_view.showGridLines is False, navn


# --- fasit 2024 -----------------------------------------------------------------------

FORVENTEDE_CELLER = (
    [("Hovedpost.", f"D{r}") for r in range(3, 20)]
    + [("Hovedpost.", f"I{r}") for r in range(3, 20)]
    + [("Resultatkomp.", c) for c in ("C15", "C16", "C17")]
    + [("Sammenligning", f"{k}{r}") for r in range(4, 13) for k in "BCDE"]
    + [("NyeSatser", f"C{r}") for r in range(5, 9)]
)


# Rader uten tall i arbeidsboken: D6/I6 og D16/I16 er tomme (etikett- og luftrader).
TOMME = [("Hovedpost.", "D6"), ("Hovedpost.", "D16"),
         ("Hovedpost.", "I6"), ("Hovedpost.", "I16")]


def test_fasit_har_alle_cellene(fasit):
    funnet = {(c["ark"], c["celle"]) for c in fasit["celler"]}
    mangler = [c for c in FORVENTEDE_CELLER if c not in funnet]
    assert sorted(mangler) == sorted(TOMME), mangler
    assert len(fasit["celler"]) == len(FORVENTEDE_CELLER) - len(TOMME)


def test_fasit_realvekst(fasit):
    r = {p["navn"]: p for p in fasit["realvekst"]}
    assert r["metode a uten RNB"]["verdi_prosent"] == 1.89
    assert r["metode a uten RNB"]["kildecelle"] == "Hovedpost.!D19"
    assert r["metode b uten RNB"]["verdi_prosent"] == -0.29
    assert r["metode a med RNB"]["verdi_prosent"] == 1.75
    assert r["metode b med RNB"]["verdi_prosent"] == -0.29
    assert r["metode b med RNB"]["kildecelle"] == "Hovedpost.!I19"


def test_fasit_kildeverdier(fasit):
    celler = {(c["ark"], c["celle"]): c for c in fasit["celler"]}
    assert celler[("Hovedpost.", "D3")]["verdi"] == 3806533
    assert celler[("Hovedpost.", "D5")]["verdi"] == 250434
    assert celler[("Hovedpost.", "D15")]["verdi"] == 4045822
    assert celler[("Hovedpost.", "D17")]["verdi"] == 239289
    assert celler[("Hovedpost.", "I4")]["verdi"] == 3892733
    assert celler[("Resultatkomp.", "C17")]["verdi"] == -5085
    assert celler[("Sammenligning", "D12")]["verdi"] == 4045822
    assert celler[("Hovedpost.", "D19")]["formel"] == "=D18-0.044"


def test_fasit_tillatte_avvik(fasit):
    ider = [a["id"] for a in fasit["tillatte_avvik"]]
    assert ider == ["A1", "A2", "A3", "A4"]
    tekst = " ".join(a["tekst"] for a in fasit["tillatte_avvik"])
    assert "Ukraina" in tekst and "USN" in tekst and "DBH" in tekst and "3,0" in tekst
