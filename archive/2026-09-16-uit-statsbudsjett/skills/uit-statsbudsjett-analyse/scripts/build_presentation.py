"""Bygger UiT-presentasjon fra ferdig vurdert innhold, med kilder i notatene."""

import argparse
import importlib.util
import json
from pathlib import Path
import subprocess
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches, Pt


GENERATOR = Path(__file__).resolve().parents[3] / 'reporting/presentation/uit-deck-generator/scripts/generate_uit_deck.py'


def clean_template(template, destination):
    """Beholder UiT-mastere og layout; fjerner gamle lysark med relasjoner."""
    presentation = Presentation(template)
    for slide_id in list(presentation.slides._sldIdLst):
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    presentation.core_properties.author = ''
    presentation.core_properties.last_modified_by = ''
    presentation.core_properties.title = ''
    presentation.core_properties.subject = ''
    presentation.core_properties.comments = ''
    presentation.save(destination)


def format_slide(slide, item):
    if item.get('type') != 'bullets':
        return
    title = slide.shapes.title
    # Sett hele geometrien: bare top/height kan nullstille arvet bredde i plassholdere.
    title.left, title.top = Inches(0.85), Inches(0.5)
    title.width, title.height = Inches(11.65), Inches(1.05)
    for paragraph in title.text_frame.paragraphs:
        paragraph.font.size = Pt(30)
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.text and shape.text == item.get('message'):
            shape.top, shape.height = Inches(1.65), Inches(0.5)
        paragraphs = [p for p in shape.text_frame.paragraphs if p.text]
        if item.get('bullets') and [p.text for p in paragraphs] == item['bullets']:
            shape.top, shape.height = Inches(2.3), Inches(4.25)
            for paragraph in paragraphs:
                paragraph.font.size = Pt(24)
                paragraph.font.name = 'Open Sans'
                paragraph.font.bold = False


def build_presentation(data, template, output):
    output = Path(output)
    if output.resolve() == Path(template).resolve():
        raise ValueError('Presentasjonen kan ikke overskrive malen.')
    module_spec = importlib.util.spec_from_file_location('uit_deck_generator', GENERATOR)
    generator = importlib.util.module_from_spec(module_spec)
    module_spec.loader.exec_module(generator)
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix='uit-deck-') as temporary:
        cleaned = Path(temporary) / 'template.pptx'
        clean_template(template, cleaned)
        generated = Path(temporary) / 'generated.pptx'
        # Generatorens ZIP-opprydding bruker /tmp; hold mellomfilen på samme filsystem.
        generator.build_deck(data, cleaned, generated, None)
        presentation = Presentation(generated)
    slide_specs = data['slides']
    if not slide_specs or slide_specs[0]['type'] != 'cover':
        slide_specs = [{'notes': data.get('outcome', '')}] + slide_specs
    if len(slide_specs) != len(presentation.slides):
        raise ValueError('Lysarkantallet samsvarer ikke med spesifikasjonen.')
    for slide, item in zip(presentation.slides, slide_specs):
        format_slide(slide, item)
        notes = item.get('notes', '')
        if item.get('sources'):
            notes += '\n\nKilder:\n' + '\n'.join(item['sources'])
        slide.notes_slide.notes_text_frame.text = notes
    presentation.core_properties.title = data['title']
    presentation.save(output)
    return generator.validate_deck(output)


def render_preview(deck, libreoffice):
    """Eksporterer den faktiske PPTX-filen til PDF og sidebilder for visuell kontroll."""
    import pymupdf
    from PIL import Image, ImageDraw

    deck = Path(deck).resolve()
    preview = deck.parent / 'preview'
    preview.mkdir(exist_ok=True)
    with TemporaryDirectory(prefix='uit-lo-profile-') as temporary:
        command = [str(libreoffice), '--headless', '--nologo', '--nodefault',
                   f'-env:UserInstallation={Path(temporary).as_uri()}',
                   '--convert-to', 'pdf', '--outdir', str(preview), str(deck)]
        subprocess.run(command, check=True, capture_output=True, text=True, timeout=120)
    pdf_path = preview / f'{deck.stem}.pdf'
    with pymupdf.open(pdf_path) as pdf:
        if len(pdf) != len(Presentation(deck).slides):
            raise ValueError('PDF-sideantall samsvarer ikke med lysarkantallet.')
        overview = Image.new('RGB', (1440, ((len(pdf) + 3) // 4) * 230), '#ddd')
        drawing = ImageDraw.Draw(overview)
        for number, page in enumerate(pdf, 1):
            path = preview / f'lysark-{number:02}.png'
            page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(path)
            with Image.open(path) as image:
                image.thumbnail((350, 197))
                x, y = ((number - 1) % 4) * 360, ((number - 1) // 4) * 230
                overview.paste(image, (x, y + 22))
                drawing.text((x + 4, y + 4), str(number), fill='black')
        overview.save(preview / 'oversikt.png')
    return {'pdf': str(pdf_path), 'pages': len(Presentation(deck).slides),
            'visual_review': 'må utføres på de genererte bildene'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec', type=Path)
    parser.add_argument('--template', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--libreoffice', type=Path, help='Render faktisk PPTX til PDF og bilder for kontroll.')
    args = parser.parse_args()
    data = json.loads(args.spec.read_text(encoding='utf-8'))
    result = build_presentation(data, args.template, args.output)
    if args.libreoffice:
        result['preview'] = render_preview(args.output, args.libreoffice)
    print(json.dumps(result, ensure_ascii=False))
