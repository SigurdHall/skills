import importlib.util
from pathlib import Path
import zipfile
import pytest
import pymupdf

from pptx import Presentation
from pptx.util import Inches


SCRIPT = Path(__file__).parents[1] / 'skills/knowledge-management/uit-statsbudsjett-analyse/scripts/build_presentation.py'
spec = importlib.util.spec_from_file_location('budget_presentation', SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def test_removes_old_slides_and_preserves_new_notes_and_sources(tmp_path):
    template = tmp_path / 'template.pptx'
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.slide_layouts[0].name = 'Tittellysbilde bred mørk'
    prs.slide_layouts[1].name = 'Tittel og innhold hvit'
    old = prs.slides.add_slide(prs.slide_layouts[1])
    old.shapes.title.text = 'PRIVATE_OLD_CONTENT'
    old.notes_slide.notes_text_frame.text = 'PRIVATE_OLD_NOTES'
    prs.save(template)
    data = {'title': 'Forslag 2024', 'slides': [
        {'type': 'cover', 'title': 'Forslag 2024', 'notes': 'Utviklingsproeve'},
        {'type': 'bullets', 'title': 'Ramme', 'bullets': ['7,3 mill.'],
         'notes': 'Nivaa, ikke oekning', 'sources': ['https://example.org/primary.pdf#page=3']},
    ]}
    output = tmp_path / 'output.pptx'
    result = module.build_presentation(data, template, output)
    assert result['slides'] == 2
    actual = Presentation(output)
    assert actual.slides[1].shapes.title.width > 0
    assert 'Nivaa, ikke oekning' in actual.slides[1].notes_slide.notes_text_frame.text
    assert 'https://example.org/primary.pdf#page=3' in actual.slides[1].notes_slide.notes_text_frame.text
    with zipfile.ZipFile(output) as archive:
        xml = b''.join(archive.read(name) for name in archive.namelist() if name.endswith('.xml'))
    assert b'PRIVATE_OLD_CONTENT' not in xml
    assert b'PRIVATE_OLD_NOTES' not in xml


def test_renderer_rejects_pdf_with_missing_slide(tmp_path, monkeypatch):
    deck = tmp_path / 'deck.pptx'
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(deck)

    def fake_converter(command, **kwargs):
        output_dir = Path(command[command.index('--outdir') + 1])
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.save(output_dir / 'deck.pdf')

    monkeypatch.setattr(module.subprocess, 'run', fake_converter)
    with pytest.raises(ValueError, match='sideantall'):
        module.render_preview(deck, 'soffice')
