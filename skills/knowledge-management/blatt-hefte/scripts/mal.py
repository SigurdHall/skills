"""Bygg arbeidsbokmalen assets/mal-rammeark.xlsx fra assets/mal-rammeark.json.

Spesifikasjonen beskriver ark, etiketter, formeltekster, tallformater og kolonnebredder.
Denne modulen legger ut layouten og skriver arbeidsboken. bygg_rammeark.py (fase 4) kaller
bygg_mal med faktisk radantall og fyller så cellene som _meta-arket peker på.

Bruk:
    mal.py --spec assets/mal-rammeark.json --output assets/mal-rammeark.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

ASSETS = Path(__file__).resolve().parent.parent / "assets"
STANDARD_SPEC = ASSETS / "mal-rammeark.json"
STANDARD_UT = ASSETS / "mal-rammeark.xlsx"


class Layout:
    """Adressene malen får: navngitte celler, radmerker per blokk og celleoppgaver."""

    def __init__(self) -> None:
        self.adresser: dict[str, tuple[str, str]] = {}
        self.radmerker: dict[str, int] = {}
        self.oppgaver: list[dict] = []

    def navngi(self, navn: str, ark: str, celle: str) -> None:
        self.adresser[navn] = (ark, celle)

    def skriv(self, ark: str, celle: str, **felt) -> None:
        self.oppgaver.append({"ark": ark, "celle": celle, **felt})

    def referanse(self, navn: str) -> str:
        ark, celle = self.adresser[navn]
        kol, rad = re.match(r"([A-Z]+)(\d+)", celle).groups()
        return f"'{ark}'!${kol}${rad}"


def _antall_rader(blokk: dict, radantall: dict[str, int] | None) -> int:
    if radantall and blokk["navn"] in radantall:
        return radantall[blokk["navn"]]
    return blokk.get("mal_radantall", 0)


def _rader(blokk: dict, ark: dict, layout: Layout, rad: int) -> int:
    navn = ark["navn"]
    if "overskrift" in blokk:
        for i, tekst in enumerate(blokk["overskrift"]):
            layout.skriv(navn, f"{chr(ord('B') + i)}{rad}", verdi=tekst, stil="overskrift")
        rad += 1
    layout.radmerker[f"{blokk['navn']}_start"] = rad
    for r in blokk["rader"]:
        layout.skriv(navn, f"B{rad}", verdi=r["etikett"], stil="fet" if r.get("inndata") else None)
        if r.get("etikett_navn"):
            layout.navngi(r["etikett_navn"], navn, f"B{rad}")
        layout.skriv(navn, f"C{rad}", formel=r.get("formel"), format=r["format"],
                     stil="inndata" if r.get("inndata") else None)
        layout.navngi(r["navn"], navn, f"C{rad}")
        rad += 1
    layout.radmerker[f"{blokk['navn']}_slutt"] = rad - 1
    return rad


def _tabell(blokk: dict, ark: dict, layout: Layout, rad: int, radantall) -> int:
    navn = ark["navn"]
    if blokk.get("blokktittel"):
        layout.skriv(navn, f"B{rad}", verdi=blokk["blokktittel"], stil="fet")
        rad += 1
    for k in blokk["kolonner"]:
        layout.skriv(navn, f"{k['kolonne']}{rad}", verdi=k["overskrift"], stil="overskrift")
    layout.navngi(f"{blokk['navn']}_overskrift", navn, f"{blokk['kolonner'][0]['kolonne']}{rad}")
    rad += 1

    start = rad
    for _ in range(_antall_rader(blokk, radantall)):
        for k in blokk["kolonner"]:
            formel = k.get("formel")
            layout.skriv(navn, f"{k['kolonne']}{rad}",
                         formel=formel.replace("{rad}", str(rad)) if formel else None,
                         format=k["format"])
        rad += 1
    slutt = rad - 1
    merker = {"start": start, "slutt": slutt, "nest_siste": max(start, slutt - 1)}
    layout.navngi(f"{blokk['navn']}_start", navn, f"{blokk['kolonner'][0]['kolonne']}{start}")

    for nokkel, merkenavn in (("sumrad", "sum"), ("kontrollrad", "kontroll")):
        spec = blokk.get(nokkel)
        if not spec:
            continue
        fet = bool(spec.get("fet"))
        layout.skriv(navn, f"B{rad}", verdi=spec["etikett"], stil="fet" if fet else None)
        layout.skriv(navn, f"{spec['kolonne']}{rad}", formel=_lokalt(spec["formel"], merker),
                     format=spec["format"], stil="fet" if fet else None)
        layout.navngi(spec["navn"], navn, f"{spec['kolonne']}{rad}")
        merker[merkenavn] = rad
        rad += 1

    for merke, verdi in merker.items():
        layout.radmerker[f"{blokk['navn']}_{merke}"] = verdi
    return rad


def _lokalt(formel: str, merker: dict[str, int]) -> str:
    """Bytt blokklokale plassholdere ({start}, {slutt}, {nest_siste}, {sum}) mot radnummer."""
    for merke, verdi in merker.items():
        formel = formel.replace("{" + merke + "}", str(verdi))
    return formel


def _fritabeller(blokk: dict, ark: dict, layout: Layout, rad: int, radantall) -> int:
    navn = ark["navn"]
    layout.skriv(navn, f"B{rad}", verdi=blokk["blokktittel"], stil="overskrift")
    for kol in blokk["kolonner"][1:]:
        layout.skriv(navn, f"{kol}{rad}", verdi=None, stil="overskrift")
    rad += 1
    start = rad
    layout.skriv(navn, f"B{rad}", verdi=blokk["tomtekst"])
    layout.navngi(f"{blokk['navn']}_tomtekst", navn, f"B{rad}")
    layout.navngi(f"{blokk['navn']}_start", navn, f"B{start}")
    rad += max(1, _antall_rader(blokk, radantall))
    layout.radmerker[f"{blokk['navn']}_start"] = start
    layout.radmerker[f"{blokk['navn']}_slutt"] = rad - 1
    return rad


def legg_ut(spec: dict, radantall: dict[str, int] | None = None) -> Layout:
    """Regn ut hvor alt havner og løs alle plassholdere i formlene."""
    layout = Layout()
    for ark in spec["ark"]:
        navn = ark["navn"]
        layout.skriv(navn, f"B{spec['layoutregler']['tittelrad']}", verdi=ark["tittel"], stil="tittel")
        if ark.get("vis_enhet"):
            layout.skriv(navn, f"B{spec['layoutregler']['enhetsrad']}",
                         verdi=spec["enhetstekst"], stil="enhet")
        rad = ark["startrad"]
        for blokk in ark["blokker"]:
            type_ = blokk["type"]
            if type_ == "rader":
                rad = _rader(blokk, ark, layout, rad)
            elif type_ == "tabell":
                rad = _tabell(blokk, ark, layout, rad, radantall)
            elif type_ == "fritabeller":
                rad = _fritabeller(blokk, ark, layout, rad, radantall)
            elif type_ == "fritekst":
                layout.skriv(navn, f"B{rad}", verdi=blokk["tekst"], stil="enhet")
                layout.navngi(blokk["navn"], navn, f"B{rad}")
                rad += 1
            else:
                raise ValueError(f"ukjent blokktype {type_!r}")
            rad += spec["layoutregler"]["tomme_rader_etter_blokk"]
    _resolver(layout)
    return layout


def _resolver(layout: Layout) -> None:
    """Bytt {blokk_merke} mot radnummer og @navn mot absolutt adresse."""
    for oppgave in layout.oppgaver:
        formel = oppgave.get("formel")
        if not formel:
            continue
        formel = re.sub(r"\{([a-z_æøå0-9]+)\}",
                        lambda m: str(layout.radmerker[m.group(1)]), formel)
        formel = re.sub(r"@([a-zæøå][a-z_æøå0-9]*)",
                        lambda m: layout.referanse(m.group(1)), formel)
        oppgave["formel"] = formel


def bygg_mal(spec: dict, output: Path, radantall: dict[str, int] | None = None) -> None:
    """Bygg arbeidsboken fra spesifikasjonen og skriv den til output."""
    layout = legg_ut(spec, radantall)
    stil = spec["stil"]
    fyll = PatternFill("solid", fgColor=stil["overskrift_bakgrunn"])
    hvit = Font(bold=True, color=stil["overskrift_tekst"])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ark_objekter = {}
    for ark in spec["ark"]:
        ws = wb.create_sheet(ark["navn"])
        ws.sheet_view.showGridLines = stil["vis_rutenett"]
        ws.freeze_panes = stil["frys_rute"]
        for kol, bredde in ark["kolonnebredder"].items():
            ws.column_dimensions[kol].width = bredde
        ark_objekter[ark["navn"]] = ws

    for o in layout.oppgaver:
        celle = ark_objekter[o["ark"]][o["celle"]]
        if o.get("formel"):
            celle.value = o["formel"]
        elif o.get("verdi") is not None:
            celle.value = o["verdi"]
        if o.get("format"):
            celle.number_format = spec["tallformat"][o["format"]]
        merke = o.get("stil")
        if merke == "overskrift":
            celle.fill, celle.font = fyll, hvit
            celle.alignment = Alignment(wrap_text=True, vertical="center")
        elif merke == "tittel":
            celle.font = Font(bold=stil["tittel"]["fet"], size=stil["tittel"]["storrelse"])
        elif merke == "enhet":
            celle.font = Font(italic=stil["enhet"]["kursiv"], size=stil["enhet"]["storrelse"])
        elif merke in ("fet", "inndata"):
            celle.font = Font(bold=True)

    meta = wb.create_sheet("_meta")
    meta.sheet_state = "hidden"
    meta.append(["type", "navn", "ark", "celle"])
    for navn in sorted(layout.adresser):
        ark, celle = layout.adresser[navn]
        meta.append(["celle", navn, ark, celle])
        wb.defined_names[navn] = DefinedName(navn, attr_text=layout.referanse(navn))
    for navn in sorted(layout.radmerker):
        meta.append(["rad", navn, "", layout.radmerker[navn]])

    wb.save(output)


def last_spec(sti: Path = STANDARD_SPEC) -> dict:
    return json.loads(sti.read_text(encoding="utf-8"))


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--spec", type=Path, default=STANDARD_SPEC)
    p.add_argument("--output", type=Path, default=STANDARD_UT)
    a = p.parse_args(argv)
    bygg_mal(last_spec(a.spec), a.output)
    print(f"mal skrevet til {a.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
