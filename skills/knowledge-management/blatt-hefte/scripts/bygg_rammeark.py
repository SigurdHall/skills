"""Fyll rammearket fra et kontraktdokument (les_blaatt_hefte.py) og malspesifikasjonen.

    bygg_rammeark.py --data blaatt-hefte-2026-forslag.json --output uit-ramme-2026-forslag.xlsx [--rnb N] [--finn-rapport finn.json]
    bygg_rammeark.py --init-template [--spec assets/mal-rammeark.json] [--output assets/mal-rammeark.xlsx]

Layouten kommer fra mal.legg_ut med faktisk radantall (heftets kolonner, indikatorer,
institusjoner), så arket for 2024 og 2026 får ulikt antall rader men samme oppbygning.
Formler skrives av malen; denne modulen skriver bare verdier og etiketter.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

HER = Path(__file__).resolve().parent
sys.path.insert(0, str(HER))
import mal  # noqa: E402

VERSJON = "1.0"


def uit_rad(dok: dict) -> dict:
    return next(r for r in dok["hovedtabell"]["rader"] if r["kortkode_norm"] == "UIT")


def avledet(dok: dict, finn_rapport: dict | None) -> dict:
    """Verdier malen ber om som ikke står direkte i kontrakten."""
    h, p = dok["hovedtabell"], dok["prisjustering"]
    uit = uit_rad(dok)
    k = h["kontroll"]
    b_ok = k["antall_band"] == k["antall_overskrifter"] and k["utgangspunkt_gjenkjent"] and k["sum_gjenkjent"]
    return {
        "uit.utgangspunkt": uit["verdier"][0],
        "uit.forslag": uit["verdier"][h["sum_indeks"]],
        "prisjustering.sats_prosent_desimal": None if p["sats_prosent"] is None else p["sats_prosent"] / 100,
        "prisjustering.ren_sats_tekst": {True: "ja", False: "nei, kolonnen inneholder mer enn satsen", None: "ikke vurdert"}[p["ren_sats"]],
        "hovedtabell.kontroll.b_tekst": f"ok ({k['antall_band']} kolonner)" if b_ok else "FEIL",
        "hovedtabell.kontroll.alle_rader_summerer": "ja" if k["alle_rader_summerer"] else "NEI",
        "finn.dom": (finn_rapport or {}).get("dom", "ikke kjørt (lokal fil)"),
    }


def slaa_opp(dok: dict, felt: str, avledede: dict):
    """Hent «a.b[0].c» fra dokumentet, eller en avledet verdi."""
    if felt in avledede:
        return avledede[felt]
    node = dok
    for del_ in felt.replace("]", "").replace("[", ".").split("."):
        if node is None:
            return None
        if del_.lstrip("-").isdigit():
            node = node[int(del_)]
        else:
            node = node.get(del_)
    return node


def institusjonsrader(dok: dict) -> tuple[list[dict], list[dict]]:
    """(universiteter, alle statlige) som rader med utgangspunkt, forslag og prisjustering."""
    h = dok["hovedtabell"]
    pris_idx = dok["prisjustering"]["kolonne_indeks"]
    navn = {i["kortkode_norm"]: i for i in dok["institusjoner"]["liste"]}
    alle = []
    for r in h["rader"]:
        inst = navn.get(r["kortkode_norm"], {})
        alle.append({
            "kortkode": r["kortkode"], "kortkode_norm": r["kortkode_norm"], "navn": inst.get("navn", ""),
            "universitet": inst.get("universitet", False),
            "utgangspunkt": r["verdier"][0], "forslag": r["verdier"][h["sum_indeks"]],
            "prisjustering": r["verdier"][pris_idx] if pris_idx is not None else None,
        })
    alle.sort(key=lambda x: x["kortkode_norm"])
    return [x for x in alle if x["universitet"]], alle


def satsblokker(dok: dict) -> tuple[list[list], list[list]]:
    """Radene til Satser-arket: [celle B, C, D, E] per rad; tom liste = tomtekst blir stående."""
    open_ = dok["satser"]["open_ramme"]
    open_rader: list[list] = []
    if open_:
        for t in open_["tabeller"]:
            open_rader.append([t["tittel"], None, None, None])
            open_rader.append(list(t["kolonner"][:4]) + [None] * (4 - len(t["kolonner"][:4])))
            for r in t["rader"]:
                open_rader.append(list(r[:4]) + [None] * (4 - len(r[:4])))
            open_rader.append([None] * 4)
    lukket = dok["satser"]["lukket_ramme"]
    lukket_rader: list[list] = []
    if lukket:
        lukket_rader.append(list(lukket["kolonner"][:4]) + [None] * (4 - len(lukket["kolonner"][:4])))
        for r in lukket["rader"]:
            lukket_rader.append(list(r[:4]) + [None] * (4 - len(r[:4])))
    return open_rader, lukket_rader


def radantall_for(dok: dict) -> dict[str, int]:
    univ, alle = institusjonsrader(dok)
    open_rader, lukket_rader = satsblokker(dok)
    return {
        "hovedposter": len(dok["hovedtabell"]["kolonner"]),
        "resultat": len(dok["resultat"]["kolonner"]) if dok["resultat"] else 1,
        "sektor_universiteter": len(univ),
        "sektor_alle": len(alle),
        "satser_open_ramme": max(1, len(open_rader)),
        "satser_lukket_ramme": max(1, len(lukket_rader)),
    }


def bygg(dok: dict, output: Path, rnb: int | None = None, finn_rapport: dict | None = None,
         spec: dict | None = None) -> dict:
    spec = spec or mal.last_spec()
    radantall = radantall_for(dok)
    layout = mal.legg_ut(spec, radantall)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    mal.bygg_mal(spec, output, radantall)

    wb = openpyxl.load_workbook(output)
    avledede = avledet(dok, finn_rapport)
    uit = uit_rad(dok)
    univ, alle = institusjonsrader(dok)
    open_rader, lukket_rader = satsblokker(dok)

    def celle(navn: str):
        ark, adr = layout.adresser[navn]
        return wb[ark][adr]

    for ark in spec["ark"]:
        ws = wb[ark["navn"]]
        for blokk in ark["blokker"]:
            if blokk["type"] == "rader":
                for r in blokk["rader"]:
                    if r.get("etikett_felt"):
                        celle(r["etikett_navn"]).value = slaa_opp(dok, r["etikett_felt"], avledede)
                    if r.get("felt"):
                        celle(r["navn"]).value = slaa_opp(dok, r["felt"], avledede)
                    if r.get("inndata") and rnb is not None:
                        celle(r["navn"]).value = rnb
            elif blokk["type"] == "tabell":
                _fyll_tabell(ws, blokk, layout, dok, uit, univ, alle)
            elif blokk["type"] == "fritabeller":
                rader = open_rader if blokk["navn"] == "satser_open_ramme" else lukket_rader
                _fyll_fritabell(ws, blokk, layout, rader)

    # RNB-etiketten får riktig år
    celle("hovedtall_rnb").offset(column=-1).value = (
        f"Tillegg i RNB {dok['kilde']['budsjettaar'] - 1}, kap. 260 post 50, hentes fra RNB-proposisjonen; tom = ikke justert")
    wb.save(output)
    return {"output": str(output), "radantall": radantall, "rnb": rnb}


def _fyll_tabell(ws, blokk: dict, layout: mal.Layout, dok: dict, uit: dict, univ: list, alle: list) -> None:
    start = layout.radmerker[f"{blokk['navn']}_start"]
    kilde = blokk["radkilde"]
    if kilde == "hovedtabell.kolonner":
        rader = [{"kolonnenavn": k, "uit_verdi": v} for k, v in zip(dok["hovedtabell"]["kolonner"], uit["verdier"])]
    elif kilde == "resultat.kolonner":
        res = dok["resultat"]
        if res is None:
            rader = [{"kolonnenavn": "Resultattabell ikke funnet i heftet", "uit_verdi": None, "er_sumkolonne": ""}]
        else:
            uit_res = next((r for r in res["rader"] if r["kortkode_norm"] == "UIT"), None)
            verdier = uit_res["verdier"] if uit_res else [None] * len(res["kolonner"])
            rader = [{"kolonnenavn": k, "uit_verdi": v, "er_sumkolonne": "ja" if i in res["sumkolonner"] else "nei"}
                     for i, (k, v) in enumerate(zip(res["kolonner"], verdier))]
    elif kilde == "institusjoner.universiteter":
        rader = univ
    elif kilde == "institusjoner.statlige":
        rader = alle
    else:
        raise ValueError(f"ukjent radkilde {kilde!r}")

    uthev = blokk.get("uthev_kortkode")
    for i, rad in enumerate(rader):
        r = start + i
        for k in blokk["kolonner"]:
            if k.get("felt"):
                ws[f"{k['kolonne']}{r}"].value = rad.get(k["felt"])
        if uthev and rad.get("kortkode_norm") == uthev:
            for k in blokk["kolonner"]:
                ws[f"{k['kolonne']}{r}"].font = Font(bold=True)


def _fyll_fritabell(ws, blokk: dict, layout: mal.Layout, rader: list[list]) -> None:
    if not rader:
        return  # tomteksten fra malen blir stående
    start = layout.radmerker[f"{blokk['navn']}_start"]
    for i, rad in enumerate(rader):
        for kol, verdi in zip(blokk["kolonner"], rad):
            ws[f"{kol}{start + i}"].value = verdi


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("--data", type=Path, help="kontraktdokument fra les_blaatt_hefte.py")
    p.add_argument("--spec", type=Path, default=mal.STANDARD_SPEC)
    p.add_argument("--output", type=Path)
    p.add_argument("--rnb", type=int)
    p.add_argument("--finn-rapport", type=Path)
    p.add_argument("--init-template", action="store_true", help="bygg bare malen fra spesifikasjonen")
    a = p.parse_args(argv)
    spec = mal.last_spec(a.spec)
    if a.init_template:
        mal.bygg_mal(spec, a.output or mal.STANDARD_UT)
        print(f"mal skrevet til {a.output or mal.STANDARD_UT}")
        return 0
    if not a.data or not a.output:
        p.error("--data og --output kreves")
    dok = json.loads(a.data.read_text(encoding="utf-8"))
    finn = json.loads(a.finn_rapport.read_text(encoding="utf-8")) if a.finn_rapport else None
    resultat = bygg(dok, a.output, rnb=a.rnb, finn_rapport=finn, spec=spec)
    print(json.dumps(resultat, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
