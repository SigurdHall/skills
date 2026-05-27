from __future__ import annotations

import argparse
import json
import sys
import warnings
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt


SKILL_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_LOGO_DIR = SKILL_ROOT / "assets" / "uit-logo-bokmal"

COLORS = {
    "dark_blue": RGBColor(0x00, 0x33, 0x49),
    "blue": RGBColor(0x00, 0x73, 0x96),
    "light_blue": RGBColor(0x59, 0xBE, 0xC9),
    "pale_blue": RGBColor(0xD8, 0xEF, 0xF7),
    "red": RGBColor(0xCB, 0x33, 0x3B),
    "yellow": RGBColor(0xF2, 0xA9, 0x00),
    "grey": RGBColor(0xF2, 0xF4, 0xF5),
    "line": RGBColor(0xD0, 0xD6, 0xDA),
    "text": RGBColor(0x1F, 0x2A, 0x33),
    "muted": RGBColor(0x58, 0x66, 0x6F),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
}

SCHEMA_EXAMPLE = {
    "deck_id": "Deck A",
    "title": "Fra rapportmigrering til styringsdata",
    "subtitle": "Hoveddeck til ledermote",
    "outcome": "Fa beslutning om pilot for styrt dataplattform.",
    "footer": "UiT virksomhetsstyring",
    "slides": [
        {
            "type": "bullets",
            "title": "Beslutningen vi ma ta",
            "message": "Dette er et styringsvalg, ikke et rent verktoyvalg.",
            "bullets": [
                "A. Isolert migrering videreforer dagens fragmentering.",
                "B. Pilot for styrt dataplattform bygger varig styringsevne.",
            ],
        }
        ,
        {
            "type": "process",
            "title": "Dataflyten",
            "message": "Fra kildesystem til lederrapport er det flere styringspunkter.",
            "steps": [
                {"label": "Kilder", "body": "Unit4/BOTT, HR og lokale filer."},
                {"label": "Plattform", "body": "Innhenting, historikk og kvalitet."},
                {"label": "Modell", "body": "Felles begreper og tilgang."},
                {"label": "Rapport", "body": "Lederflate og analyse."}
            ]
        }
    ],
}


def remove_all_slides(prs: Presentation) -> None:
    for slide_id in list(prs.slides._sldIdLst):  # noqa: SLF001
        prs.slides._sldIdLst.remove(slide_id)  # noqa: SLF001


def remove_empty_placeholders(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in list(slide.shapes):
            if shape.is_placeholder and not getattr(shape, "text", "").strip():
                shape.element.getparent().remove(shape.element)


def add_click_appear_animations(slide, shapes) -> None:
    shape_ids = [str(shape.shape_id) for shape in shapes if shape is not None]
    if not shape_ids:
        return
    slide_el = slide._element  # noqa: SLF001 - python-pptx has no public animation API.
    for timing in list(slide_el.findall(qn("p:timing"))):
        slide_el.remove(timing)
    child_xml = []
    bld_xml = []
    next_id = 3
    for spid in shape_ids:
        child_xml.append(
            f"""
            <p:par><p:cTn id="{next_id}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">
              <p:stCondLst><p:cond delay="0"/></p:stCondLst>
              <p:childTnLst><p:set><p:cBhvr><p:cTn id="{next_id + 1}" dur="1" fill="hold">
                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
              </p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
              <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
              </p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst>
            </p:cTn></p:par>
            """
        )
        bld_xml.append(f'<p:bldP spid="{spid}" grpId="0" animBg="1"/>')
        next_id += 2
    timing_xml = f"""
    <p:timing {nsdecls("p")}>
      <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">
          <p:childTnLst>{''.join(child_xml)}</p:childTnLst>
        </p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
        <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst>
      </p:cTn></p:par></p:tnLst>
      <p:bldLst>{''.join(bld_xml)}</p:bldLst>
    </p:timing>
    """
    timing_el = parse_xml(timing_xml)
    ext_lst = slide_el.find(qn("p:extLst"))
    if ext_lst is not None:
        slide_el.insert(list(slide_el).index(ext_lst), timing_el)
    else:
        slide_el.append(timing_el)


def add_bullet_format(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001
    p_pr.set("marL", "342900")
    p_pr.set("indent", "-171450")
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip"):
        for child in list(p_pr.findall(qn(tag))):
            p_pr.remove(child)
    p_pr.insert(0, parse_xml(f'<a:buChar {nsdecls("a")} char="•"/>'))
    p_pr.insert(0, parse_xml(f'<a:buFont {nsdecls("a")} typeface="+mj-lt"/>'))


def dedupe_pptx_zip(path: Path) -> None:
    with ZipFile(path, "r") as source:
        last_members: dict[str, bytes] = {}
        order: list[str] = []
        for info in source.infolist():
            if info.filename not in last_members:
                order.append(info.filename)
            last_members[info.filename] = source.read(info.filename)

    with NamedTemporaryFile(delete=False, suffix=".pptx") as temp_file:
        temp_path = Path(temp_file.name)

    with ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
        for filename in order:
            target.writestr(filename, last_members[filename])

    temp_path.replace(path)


def find_layout(prs: Presentation, names: list[str], fallback_index: int = 0):
    for name in names:
        for slide_layout in prs.slide_layouts:
            if slide_layout.name == name:
                return slide_layout
    return prs.slide_layouts[fallback_index]


def logo_path(explicit: str | None = None) -> Path | None:
    if explicit:
        path = Path(explicit)
        return path if path.exists() else None
    matches = list(DEFAULT_LOGO_DIR.rglob("UiT_Logo_Bok_Bla_RGB.png"))
    return matches[0] if matches else None


def add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int = 18,
    bold: bool = False,
    color: RGBColor | None = None,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color or COLORS["text"]
    return box


def add_rect(slide, left: float, top: float, width: float, height: float, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def set_title(slide, text: str, size: int = 30) -> None:
    if slide.shapes.title:
        title = slide.shapes.title
        title.text = text
        p = title.text_frame.paragraphs[0]
        p.font.size = Pt(size)
        p.font.bold = True
        p.font.color.rgb = COLORS["dark_blue"]
    else:
        add_text(slide, text, 0.9, 0.42, 11.4, 0.75, size=size, bold=True, color=COLORS["dark_blue"])


def add_footer(slide, deck_id: str, footer: str, page: int) -> None:
    text = " | ".join(part for part in [deck_id, footer, str(page)] if part)
    add_text(slide, text, 0.65, 7.08, 11.9, 0.22, size=8, color=COLORS["muted"], align=PP_ALIGN.RIGHT)


def add_logo(slide, path: Path | None, cover: bool = False) -> None:
    if not path:
        return
    left, top, width = (8.75, 6.42, 2.9) if cover else (10.65, 0.22, 1.35)
    try:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    except Exception:
        return


def add_bullets(slide, bullets: list[str], left: float, top: float, width: float, height: float, size: int) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        add_bullet_format(p)
        p.font.size = Pt(size)
        p.font.color.rgb = COLORS["text"]
        p.space_after = Pt(12)
        p.line_spacing = 1.08
    return box


def build_cover(prs: Presentation, spec: dict, logo: Path | None) -> None:
    slide = prs.slides.add_slide(
        find_layout(prs, ["Tittellysbilde bred m\u00f8rk", "Tittellysbilde bred mørk"], 0)
    )
    if slide.shapes.title:
        slide.shapes.title.text = spec.get("title", "UiT-presentasjon")
        p = slide.shapes.title.text_frame.paragraphs[0]
        p.font.size = Pt(34)
        p.font.bold = True
        p.font.color.rgb = COLORS["white"]
    else:
        add_text(slide, spec.get("title", "UiT-presentasjon"), 0.75, 1.1, 5.0, 1.2, 34, True, COLORS["white"])

    add_rect(slide, 7.0, 0, 6.4, 7.5, COLORS["pale_blue"])
    add_rect(slide, 7.0, 0, 0.14, 7.5, COLORS["yellow"])
    add_text(slide, spec.get("deck_id", ""), 7.55, 1.0, 4.6, 0.4, 18, True, COLORS["dark_blue"])
    add_text(slide, spec.get("subtitle", ""), 0.78, 3.15, 4.6, 0.55, 17, False, COLORS["white"])
    add_text(slide, "Mål for møtet", 7.55, 2.05, 4.6, 0.35, 14, True, COLORS["dark_blue"])
    outcome_shape = add_text(slide, spec.get("outcome", ""), 7.55, 2.48, 4.7, 2.2, 19, False, COLORS["text"])
    add_logo(slide, logo, cover=True)
    add_click_appear_animations(slide, [outcome_shape])


def build_bullets(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.45, 16, True, COLORS["red"])]
    shapes.append(add_bullets(slide, slide_spec.get("bullets", []), 1.0, 2.0, 11.1, 4.45, 20))
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_cards(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.45, 16, True, COLORS["red"])]
    cards = slide_spec.get("cards", [])
    columns = int(slide_spec.get("columns", 2))
    left0, top0, gap_x, gap_y = 0.85, 1.95, 0.25, 0.25
    card_w = (11.65 - gap_x * (columns - 1)) / columns
    rows = max(1, (len(cards) + columns - 1) // columns)
    card_h = min(1.28, (4.8 - gap_y * (rows - 1)) / rows)
    for i, card in enumerate(cards):
        col, row = i % columns, i // columns
        left = left0 + col * (card_w + gap_x)
        top = top0 + row * (card_h + gap_y)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(card_w), Inches(card_h))
        shape.adjustments[0] = 0.08
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["grey"]
        shape.line.color.rgb = COLORS["line"]
        shapes.append(shape)
        add_text(slide, card.get("title", ""), left + 0.15, top + 0.12, card_w - 0.3, 0.3, 14, True, COLORS["dark_blue"])
        add_text(slide, card.get("body", ""), left + 0.15, top + 0.46, card_w - 0.3, card_h - 0.52, 12)
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_two_column(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.38, 15, True, COLORS["red"])]
    shapes.append(add_rect(slide, 0.85, 1.85, 5.55, 0.12, COLORS["red"]))
    shapes.append(add_rect(slide, 6.9, 1.85, 5.55, 0.12, COLORS["blue"]))
    add_text(slide, slide_spec.get("left_title", "I dag"), 0.85, 2.04, 5.55, 0.35, 17, True, COLORS["dark_blue"])
    add_text(slide, slide_spec.get("right_title", "Målbilde"), 6.9, 2.04, 5.55, 0.35, 17, True, COLORS["dark_blue"])
    shapes.append(add_bullets(slide, slide_spec.get("left_items", []), 0.9, 2.5, 5.35, 3.8, 15))
    shapes.append(add_bullets(slide, slide_spec.get("right_items", []), 6.95, 2.5, 5.35, 3.8, 15))
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_table(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.38, 15, True, COLORS["red"])]
    headers = slide_spec.get("headers", ["", ""])
    rows = slide_spec.get("rows", [])
    x, y, col1, col2, row_h = 0.85, 1.85, 3.1, 8.45, 0.78
    shapes.append(add_rect(slide, x, y, col1 + col2, 0.44, COLORS["dark_blue"]))
    add_text(slide, headers[0], x + 0.12, y + 0.08, col1 - 0.2, 0.25, 11, True, COLORS["white"])
    add_text(slide, headers[1], x + col1 + 0.12, y + 0.08, col2 - 0.2, 0.25, 11, True, COLORS["white"])
    for i, row in enumerate(rows):
        top = y + 0.52 + i * row_h
        shapes.append(add_rect(slide, x, top, col1 + col2, row_h - 0.03, COLORS["grey"] if i % 2 == 0 else COLORS["white"]))
        add_text(slide, row[0], x + 0.12, top + 0.08, col1 - 0.2, row_h - 0.16, 11, True)
        add_text(slide, row[1], x + col1 + 0.12, top + 0.08, col2 - 0.2, row_h - 0.16, 11)
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_timeline(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.38, 15, True, COLORS["red"])]
    steps = slide_spec.get("steps", [])
    left, top, width = 0.9, 2.25, 11.5
    step_w = width / max(1, len(steps))
    shapes.append(add_rect(slide, left, top + 0.58, width, 0.07, COLORS["dark_blue"]))
    for i, step in enumerate(steps):
        cx = left + i * step_w + step_w / 2
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.28), Inches(top + 0.32), Inches(0.56), Inches(0.56))
        oval.fill.solid()
        oval.fill.fore_color.rgb = COLORS["red"] if i == 0 else COLORS["dark_blue"]
        oval.line.fill.background()
        shapes.append(oval)
        add_text(slide, step.get("label", ""), cx - step_w / 2 + 0.08, top + 1.04, step_w - 0.16, 0.35, 15, True, align=PP_ALIGN.CENTER)
        add_text(slide, step.get("body", ""), cx - step_w / 2 + 0.12, top + 1.48, step_w - 0.24, 1.45, 12, align=PP_ALIGN.CENTER)
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_quote(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec.get("title", "Hovedbudskap"))
    shapes = [add_rect(slide, 1.1, 2.0, 0.12, 2.9, COLORS["yellow"])]
    shapes.append(add_text(slide, slide_spec.get("quote", ""), 1.45, 2.0, 10.2, 2.4, 28, True, COLORS["dark_blue"]))
    shapes.append(add_text(slide, slide_spec.get("attribution", ""), 1.45, 4.55, 8.8, 0.4, 13, False, COLORS["muted"]))
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_process(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.38, 15, True, COLORS["red"])]
    steps = slide_spec.get("steps", [])
    left, top, width, gap = 0.85, 2.15, 11.55, 0.18
    step_w = (width - gap * max(0, len(steps) - 1)) / max(1, len(steps))
    for i, step in enumerate(steps):
        x = left + i * (step_w + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(step_w), Inches(1.55))
        box.adjustments[0] = 0.08
        box.fill.solid()
        box.fill.fore_color.rgb = COLORS["pale_blue"] if i % 2 == 0 else COLORS["grey"]
        box.line.color.rgb = COLORS["line"]
        shapes.append(box)
        add_text(slide, step.get("label", ""), x + 0.12, top + 0.16, step_w - 0.24, 0.28, 12, True, COLORS["dark_blue"], PP_ALIGN.CENTER)
        add_text(slide, step.get("body", ""), x + 0.12, top + 0.52, step_w - 0.24, 0.78, 10, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + step_w - 0.06), Inches(top + 0.57), Inches(0.32), Inches(0.26))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COLORS["yellow"]
            arrow.line.fill.background()
            shapes.append(arrow)
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_hub(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.38, 15, True, COLORS["red"])]
    center = slide_spec.get("center", "Felles modell")
    spokes = slide_spec.get("spokes", [])[:4]
    top = 1.95
    hub = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.2), Inches(top + 1.55), Inches(2.7), Inches(1.15))
    hub.fill.solid()
    hub.fill.fore_color.rgb = COLORS["dark_blue"]
    hub.line.fill.background()
    shapes.append(hub)
    add_text(slide, center, 5.35, top + 1.82, 2.4, 0.45, 15, True, COLORS["white"], PP_ALIGN.CENTER)
    positions = [(0.9, top), (9.25, top), (0.9, top + 3.1), (9.25, top + 3.1)]
    for item, (x, y) in zip(spokes, positions):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.25), Inches(1.25))
        card.adjustments[0] = 0.08
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS["grey"]
        card.line.color.rgb = COLORS["line"]
        shapes.append(card)
        add_text(slide, item.get("title", ""), x + 0.12, y + 0.14, 3.0, 0.28, 12, True, COLORS["dark_blue"])
        add_text(slide, item.get("body", ""), x + 0.12, y + 0.48, 3.0, 0.52, 10)
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


def build_decision(prs: Presentation, spec: dict, slide_spec: dict, page: int, logo: Path | None) -> None:
    slide = prs.slides.add_slide(find_layout(prs, ["Tittel og innhold hvit"], 9))
    set_title(slide, slide_spec["title"])
    shapes = [add_text(slide, slide_spec.get("message", ""), 0.92, 1.34, 11.2, 0.38, 15, True, COLORS["red"])]
    options = slide_spec.get("options", [])[:2]
    for idx, option in enumerate(options):
        x = 0.95 if idx == 0 else 6.75
        color = COLORS["red"] if idx == 0 else COLORS["dark_blue"]
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.2), Inches(5.25), Inches(2.25))
        card.adjustments[0] = 0.08
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS["grey"]
        card.line.color.rgb = color
        shapes.append(card)
        add_rect(slide, x, 2.2, 5.25, 0.12, color)
        add_text(slide, option.get("title", ""), x + 0.22, 2.48, 4.8, 0.42, 18, True, color)
        add_text(slide, option.get("body", ""), x + 0.22, 3.04, 4.75, 0.95, 15)
    add_logo(slide, logo)
    add_footer(slide, spec.get("deck_id", ""), spec.get("footer", ""), page)
    add_click_appear_animations(slide, shapes)


BUILDERS = {
    "bullets": build_bullets,
    "cards": build_cards,
    "two-column": build_two_column,
    "table": build_table,
    "timeline": build_timeline,
    "quote": build_quote,
    "process": build_process,
    "hub": build_hub,
    "decision": build_decision,
}


def build_deck(spec: dict, template: Path, output: Path, logo: Path | None) -> None:
    if not template.exists():
        raise FileNotFoundError(f"UiT template not found: {template}")
    prs = Presentation(str(template))
    remove_all_slides(prs)
    slides = spec.get("slides", [])
    if not slides or slides[0].get("type") != "cover":
        build_cover(prs, spec, logo)
        page_offset = 1
    else:
        cover_spec = {**spec, **slides[0]}
        build_cover(prs, cover_spec, logo)
        slides = slides[1:]
        page_offset = 1

    for i, slide_spec in enumerate(slides, start=1 + page_offset):
        slide_type = slide_spec.get("type", "bullets")
        builder = BUILDERS.get(slide_type)
        if not builder:
            raise ValueError(f"Unsupported slide type: {slide_type}")
        builder(prs, spec, slide_spec, i, logo)

    remove_empty_placeholders(prs)
    output.parent.mkdir(parents=True, exist_ok=True)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Duplicate name:*", category=UserWarning)
        prs.save(output)
    dedupe_pptx_zip(output)


def validate_deck(path: Path) -> dict:
    prs = Presentation(str(path))
    empty = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.is_placeholder and not getattr(shape, "text", "").strip():
                empty.append((slide_index, shape.name))
    with ZipFile(path) as zf:
        names = [info.filename for info in zf.infolist()]
    return {
        "slides": len(prs.slides),
        "empty_placeholders": len(empty),
        "duplicate_zip_members": len(names) - len(set(names)),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate a UiT-branded PowerPoint deck from JSON.")
    parser.add_argument("--spec", type=Path, help="Path to JSON deck spec.")
    parser.add_argument("--output", type=Path, help="Output .pptx path.")
    parser.add_argument("--template", type=Path, help="UiT .pptx/.potx template path.")
    parser.add_argument("--logo", type=str, default=None, help="Optional UiT logo PNG path.")
    parser.add_argument("--print-schema", action="store_true", help="Print a minimal JSON spec example.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.print_schema:
        print(json.dumps(SCHEMA_EXAMPLE, indent=2, ensure_ascii=False))
        return 0
    if not args.spec or not args.output or not args.template:
        print("Require --spec, --output, and --template unless --print-schema is used.", file=sys.stderr)
        return 2
    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    logo = logo_path(args.logo)
    build_deck(spec, args.template, args.output, logo)
    print(json.dumps(validate_deck(args.output), indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
